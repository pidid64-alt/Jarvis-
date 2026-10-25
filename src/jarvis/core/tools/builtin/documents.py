"""Инструменты чтения локальных документов: PDF, Word (.docx), Excel (.xlsx).

Файлы ищутся в data/docs (своего рода «рабочая папка» ассистента) и по
абсолютному пути. Всё содержимое документов — недоверенные данные.
"""

from __future__ import annotations

from pathlib import Path

from jarvis.config import Settings
from jarvis.core.tools.base import Tool, ToolResult


def _resolve_path(settings: Settings, raw: str) -> Path:
    path = Path(raw.strip().strip('"'))
    if not path.is_absolute():
        path = Path(settings.data_dir) / "docs" / path
    return path


def _pdf_text(path: Path, max_pages: int) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages[:max_pages], 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"— стр. {i} —\n{text}")
    return "\n\n".join(pages)


def _docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:  # таблицы — построчно
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx_text(path: Path, max_rows: int) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    chunks = []
    for sheet in workbook.worksheets:
        lines = [f"[Лист: {sheet.title}]"]
        for row in sheet.iter_rows(max_row=max_rows, values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells))
        if len(lines) > 1:
            chunks.append("\n".join(lines))
    workbook.close()
    return "\n\n".join(chunks)


def create_tools(settings: Settings, max_chars: int = 12000) -> list[Tool]:
    async def read_document(args: dict) -> ToolResult:
        raw = str(args.get("path", ""))
        if not raw.strip():
            return ToolResult(output="Ошибка: укажи путь или имя файла из data/docs.")
        path = _resolve_path(settings, raw)
        if not path.exists():
            docs_dir = Path(settings.data_dir) / "docs"
            return ToolResult(
                output=f"Файл не найден: {path}. Доступные в data/docs: "
                + (", ".join(p.name for p in docs_dir.glob("*") if p.is_file()) or "(пусто)")
            )
        limit = int(args.get("max_chars", max_chars))

        try:
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                text = _pdf_text(path, max_pages=int(args.get("max_pages", 30)))
            elif suffix in (".docx",):
                text = _docx_text(path)
            elif suffix in (".xlsx", ".xlsm"):
                text = _xlsx_text(path, max_rows=int(args.get("max_rows", 200)))
            else:
                text = path.read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            return ToolResult(output=f"Не удалось прочитать {path.name}: {e}")

        if not text.strip():
            return ToolResult(output=f"Документ {path.name} пуст или не содержит текста (скан без OCR?).")
        if len(text) > limit:
            text = text[:limit] + "\n[…документ обрезан]"
        return ToolResult(output=f"Документ {path.name}:\n\n{text}", untrusted=True)

    async def list_documents(args: dict) -> ToolResult:
        docs_dir = Path(settings.data_dir) / "docs"
        docs_dir.mkdir(parents=True, exist_ok=True)
        files = sorted(p for p in docs_dir.glob("*") if p.is_file())
        if not files:
            return ToolResult(output=f"В {docs_dir} файлов нет. Положи документы туда или читай по полному пути.")
        size_kb = lambda p: f"{p.stat().st_size / 1024:.0f} КБ"
        lines = [f"- {p.name} ({size_kb(p)})" for p in files[:50]]
        return ToolResult(output=f"Документы в {docs_dir}:\n" + "\n".join(lines))

    return [
        Tool(
            name="read_document",
            description=(
                "Прочитать локальный документ: PDF, Word (.docx), Excel (.xlsx) или текстовый файл. "
                "Возвращает текст и таблицы. Аргумент path — имя файла из data/docs либо полный путь."
            ),
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Имя файла из data/docs или полный путь"},
                    "max_chars": {"type": "integer", "description": "Лимит длины текста"},
                    "max_pages": {"type": "integer", "description": "PDF: сколько страниц"},
                    "max_rows": {"type": "integer", "description": "Excel: сколько строк на лист"},
                },
                "required": ["path"],
            },
            handler=read_document,
            side_effect="read",
            returns_untrusted=True,
            timeout=60.0,
        ),
        Tool(
            name="list_documents",
            description="Список документов в папке data/docs.",
            parameters={"type": "object", "properties": {}},
            handler=list_documents,
            side_effect="read",
        ),
    ]
